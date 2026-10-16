import csv
import hmac
import io
import socket
from base64 import b64encode
from base64 import b64decode
from datetime import UTC, datetime
from hashlib import sha256
from secrets import token_urlsafe
from uuid import uuid4

import qrcode
import qrcode.image.svg
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.approval.database import get_approval_db
from app.approval.models import ApprovalUser
from app.approval.serializers import serialize_approval_user
from app.business.models import Asset, AssetException, InventoryRecord, InventoryTask
from app.business.schemas import AssetCreate, AssetImportCommit, AssetUpdate, InventoryRecordUpdate, InventoryTaskCreate, InventoryTaskUpdate, MobileInventorySubmit, ResolveExceptionRequest
from app.core.config import settings
from app.core.security import require_permission
from app.db.database import get_db
from app.db.models import User
from app.integrations.storage import StorageUploadError, upload_bytes

router = APIRouter(prefix="/assets", tags=["assets"])


def inventory_scope_condition(scope: str):
    """Translate the UI scope value into a database filter."""
    scope_filters = {
        "部门：": Asset.department,
        "类别：": Asset.category,
        "位置：": Asset.location,
    }
    for prefix, column in scope_filters.items():
        if scope.startswith(prefix):
            value = scope.removeprefix(prefix).strip()
            return column == value if value else None
    return None


def asset_payload(row: Asset) -> dict:
    return {"id": row.id, "asset_uuid": row.asset_uuid, "asset_code": row.asset_code, "name": row.name, "category": row.category, "department": row.department, "owner": row.owner, "location": row.location, "status": row.status, "purchase_date": row.purchase_date, "value": row.value, "last_inventory_at": row.last_inventory_at, "created_at": row.created_at, "updated_at": row.updated_at}


def task_payload(row: InventoryTask) -> dict:
    return {"id": row.id, "name": row.name, "scope": row.scope, "owner": row.owner, "status": row.status, "due_date": row.due_date, "total_count": row.total_count, "checked_count": row.checked_count, "exception_count": row.exception_count, "created_at": row.created_at, "updated_at": row.updated_at}


def record_payload(row: InventoryRecord, db: Session, asset: Asset | None = None) -> dict:
    asset = asset or db.get(Asset, row.asset_id)
    return {
        "id": row.id,
        "task_id": row.task_id,
        "asset_id": row.asset_id,
        "asset_code": asset.asset_code if asset else "",
        "asset_name": asset.name if asset else "",
        "department": asset.department if asset else "",
        "owner": asset.owner if asset else "",
        "ledger_location": asset.location if asset else "",
        "result": row.result,
        "actual_location": row.actual_location,
        "remark": row.remark,
        "photo_url": row.photo_url,
        "checked_by": row.checked_by,
        "checked_at": row.checked_at,
    }


def active_record_payload(asset: Asset, db: Session) -> dict | None:
    row = db.scalar(
        select(InventoryRecord)
        .join(InventoryTask, InventoryRecord.task_id == InventoryTask.id)
        .where(InventoryRecord.asset_id == asset.id, InventoryTask.status == "进行中")
        .order_by(InventoryRecord.id.desc())
    )
    if row is None:
        return None
    task = db.get(InventoryTask, row.task_id)
    locked = row.result != "未盘点"
    return {
        **record_payload(row, db),
        "task_name": task.name if task else "",
        "task_scope": task.scope if task else "",
        "task_due_date": task.due_date if task else None,
        "task_status": task.status if task else "",
        "locked": locked,
    }


def persist_photo_url(photo_url: str | None, asset: Asset, record_id: int) -> str | None:
    if not photo_url or not photo_url.startswith("data:"):
        return photo_url
    try:
        header, encoded = photo_url.split(",", 1)
        content_type = header[5:].split(";", 1)[0] or "image/jpeg"
        content = b64decode(encoded, validate=True)
    except (ValueError, UnicodeError):
        raise HTTPException(status_code=400, detail="现场照片格式无效") from None
    try:
        uploaded = upload_bytes(
            f"inventory/{asset.asset_uuid}/{record_id}-{uuid4().hex}",
            content,
            content_type,
        )
    except StorageUploadError as exc:
        if settings.app_env.lower() == "production":
            raise HTTPException(status_code=503, detail="云端照片存储不可用，请稍后重试") from exc
        return photo_url
    if uploaded:
        return uploaded
    if settings.app_env.lower() == "production":
        raise HTTPException(status_code=503, detail="生产环境必须配置云端照片存储")
    return photo_url


def exception_payload(row: AssetException, db: Session, asset: Asset | None = None) -> dict:
    asset = asset or db.get(Asset, row.asset_id)
    return {"id": row.id, "asset_id": row.asset_id, "asset_code": asset.asset_code if asset else "", "asset_name": asset.name if asset else "", "department": asset.department if asset else "", "type": row.type, "description": row.description, "status": row.status, "resolution": row.resolution, "resolved_by": row.resolved_by, "resolved_at": row.resolved_at, "created_at": row.created_at}


@router.get("/summary")
def asset_summary(_: User = Depends(require_permission("asset:view")), db: Session = Depends(get_db)) -> dict:
    total, total_value, exception_count, inventory_total, inventory_checked = db.execute(
        select(
            select(func.count(Asset.id)).scalar_subquery(),
            select(func.coalesce(func.sum(Asset.value), 0)).scalar_subquery(),
            select(func.count(AssetException.id))
            .where(AssetException.status != "已归档")
            .scalar_subquery(),
            select(func.coalesce(func.sum(InventoryTask.total_count), 0)).scalar_subquery(),
            select(func.coalesce(func.sum(InventoryTask.checked_count), 0)).scalar_subquery(),
        )
    ).one()
    completion = (
        min(100, round((inventory_checked / inventory_total) * 100))
        if inventory_total
        else 0
    )
    return {"data": {"asset_count": total, "total_value": total_value, "exception_count": exception_count, "inventory_completion": completion}}


@router.get("/overview")
def asset_overview(
    _: User = Depends(require_permission("asset:view")),
    db: Session = Depends(get_db),
    approval_db: Session = Depends(get_approval_db),
) -> dict:
    summary = asset_summary(_, db)["data"]
    assets = db.scalars(select(Asset).order_by(Asset.created_at.desc())).all()
    tasks = db.scalars(
        select(InventoryTask).order_by(InventoryTask.created_at.desc())
    ).all()
    exception_rows = db.execute(
        select(AssetException, Asset)
        .join(Asset, Asset.id == AssetException.asset_id)
        .order_by(AssetException.created_at.desc())
    ).all()
    users = approval_db.scalars(
        select(ApprovalUser).order_by(ApprovalUser.display_name)
    ).all()
    return {
        "data": {
            "summary": summary,
            "assets": [asset_payload(row) for row in assets],
            "tasks": [task_payload(row) for row in tasks],
            "exceptions": [
                exception_payload(exception, db, asset) for exception, asset in exception_rows
            ],
            "inventory_owners": [serialize_approval_user(user) for user in users],
        }
    }


@router.get("")
def list_assets(keyword: str = "", status_filter: str = Query("", alias="status"), category: str = "", department: str = "", _: User = Depends(require_permission("asset:view")), db: Session = Depends(get_db)) -> dict:
    query = select(Asset).order_by(Asset.created_at.desc())
    if keyword:
        term = f"%{keyword}%"
        query = query.where(or_(Asset.asset_code.like(term), Asset.name.like(term), Asset.department.like(term), Asset.owner.like(term)))
    if status_filter: query = query.where(Asset.status == status_filter)
    if category: query = query.where(Asset.category == category)
    if department: query = query.where(Asset.department == department)
    return {"data": [asset_payload(row) for row in db.scalars(query).all()]}


def normalize_import_key(value: str) -> str:
    aliases = {
        "资产编号": "asset_code",
        "固定ID": "asset_code",
        "资产名称": "name",
        "名称": "name",
        "类别": "category",
        "所属部门": "department",
        "部门": "department",
        "责任人": "owner",
        "位置": "location",
        "状态": "status",
        "资产价值": "value",
        "价值": "value",
    }
    key = value.strip()
    return aliases.get(key, key)


def parse_asset_rows(file: UploadFile) -> list[dict]:
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    content = file.file.read()
    if suffix == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [normalize_import_key(str(value or "")) for value in rows[0]]
        return [
            {header: value for header, value in zip(headers, row, strict=False)}
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = [normalize_import_key(value) for value in (reader.fieldnames or [])]
    return [row for row in reader]


def row_to_import(row: dict, row_number: int, existing_codes: set[str], seen_codes: set[str]) -> dict:
    required = ["asset_code", "name", "category", "department", "owner", "location"]
    cleaned = {key: str(row.get(key) or "").strip() for key in required}
    cleaned["status"] = str(row.get("status") or "在用").strip()
    try:
        cleaned["value"] = float(row.get("value") or 0)
    except (TypeError, ValueError):
        cleaned["value"] = -1
    errors = [f"缺少{key}" for key in required if not cleaned[key]]
    if cleaned["value"] < 0:
        errors.append("资产价值不能为负数")
    if cleaned["asset_code"] in existing_codes:
        errors.append("资产编号已存在")
    if cleaned["asset_code"] in seen_codes:
        errors.append("文件内资产编号重复")
    seen_codes.add(cleaned["asset_code"])
    return {"row_number": row_number, **cleaned, "valid": not errors, "errors": errors}


def asset_qr_token(asset: Asset) -> str:
    return hmac.new(
        settings.jwt_secret.encode(),
        f"{asset.asset_uuid}:{asset.qr_secret}".encode(),
        sha256,
    ).hexdigest()


def qr_image_data_url(value: str) -> str:
    image_factory = qrcode.image.svg.SvgPathImage
    image = qrcode.make(value, image_factory=image_factory)
    buffer = io.BytesIO()
    image.save(buffer)
    encoded = b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/svg+xml;base64,{encoded}"


def get_lan_ip() -> str:
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sock:
            sock.connect(("8.8.8.8", 80))
            return sock.getsockname()[0]
    except OSError:
        return "127.0.0.1"


def mobile_app_base_url(request: Request) -> tuple[str, str]:
    if settings.public_app_url:
        return settings.public_app_url.rstrip("/"), ""
    origin = request.headers.get("origin", "").rstrip("/")
    if origin and "127.0.0.1" not in origin and "localhost" not in origin:
        return origin, ""
    lan_ip = get_lan_ip()
    if lan_ip != "127.0.0.1":
        return (
            f"http://{lan_ip}:5173",
            "未配置 PUBLIC_APP_URL，已使用局域网地址。手机需与电脑在同一 WiFi，且前端服务需用 0.0.0.0 启动。",
        )
    return (
        "http://127.0.0.1:5173",
        "当前二维码仍是本机地址，微信手机无法打开。请配置 PUBLIC_APP_URL 为公网或局域网可访问地址。",
    )


def qr_payload(asset: Asset, request: Request) -> dict:
    token = asset_qr_token(asset)
    base_url, warning = mobile_app_base_url(request)
    mobile_url = f"{base_url}/assets/mobile-inventory?asset_uuid={asset.asset_uuid}&token={token}"
    return {
        "asset_id": asset.id,
        "asset_uuid": asset.asset_uuid,
        "asset_code": asset.asset_code,
        "token": token,
        "mobile_url": mobile_url,
        "mobile_url_warning": warning,
        "qr_image": qr_image_data_url(mobile_url),
    }


def refresh_task_progress(task: InventoryTask, db: Session, now: datetime) -> None:
    task.checked_count = db.scalar(select(func.count(InventoryRecord.id)).where(InventoryRecord.task_id == task.id, InventoryRecord.result != "未盘点")) or 0
    task.exception_count = db.scalar(
        select(func.count(AssetException.id))
        .join(InventoryRecord, AssetException.inventory_record_id == InventoryRecord.id)
        .where(InventoryRecord.task_id == task.id, AssetException.status != "已归档")
    ) or 0
    if task.status not in {"已结束", "已完成"}:
        task.status = "已完成" if task.total_count and task.checked_count >= task.total_count else "进行中"
    task.updated_at = now


def apply_inventory_record_update(
    record: InventoryRecord,
    asset: Asset | None,
    task: InventoryTask | None,
    *,
    result: str,
    actual_location: str | None,
    remark: str | None,
    photo_url: str | None,
    checked_by: str,
    db: Session,
) -> None:
    now = datetime.now(UTC)
    record.result = result
    record.actual_location = actual_location or (asset.location if asset else None)
    record.remark = remark
    record.photo_url = photo_url
    record.checked_by = checked_by
    record.checked_at = now
    if asset:
        asset.last_inventory_at = now.date()

    existing_exception = db.scalar(select(AssetException).where(AssetException.inventory_record_id == record.id))
    if result == "正常":
        if existing_exception and existing_exception.status != "已归档":
            existing_exception.status = "已归档"
            existing_exception.resolution = "盘点结果已更正为正常"
            existing_exception.resolved_by = checked_by
            existing_exception.resolved_at = now
    elif existing_exception:
        existing_exception.type = result
        existing_exception.description = remark or f"盘点结果：{result}"
        existing_exception.status = "待核实"
    else:
        db.add(AssetException(asset_id=record.asset_id, inventory_record_id=record.id, type=result, description=remark or f"盘点结果：{result}"))

    db.flush()
    if task:
        refresh_task_progress(task, db, now)


@router.post("/import/preview")
def preview_asset_import(file: UploadFile = File(...), _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    rows = parse_asset_rows(file)
    existing_codes = set(db.scalars(select(Asset.asset_code)).all())
    seen_codes: set[str] = set()
    parsed = [row_to_import(row, index + 2, existing_codes, seen_codes) for index, row in enumerate(rows)]
    return {
        "data": {
            "total": len(parsed),
            "valid_count": sum(1 for row in parsed if row["valid"]),
            "error_count": sum(1 for row in parsed if not row["valid"]),
            "rows": parsed,
        }
    }


@router.post("/import/commit", status_code=status.HTTP_201_CREATED)
def commit_asset_import(payload: AssetImportCommit, _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    existing_codes = set(db.scalars(select(Asset.asset_code)).all())
    created: list[Asset] = []
    skipped: list[dict] = []
    for row in payload.rows:
        if row.asset_code in existing_codes:
            skipped.append({"asset_code": row.asset_code, "reason": "资产编号已存在"})
            continue
        asset = Asset(**row.model_dump(exclude={"row_number"}), asset_uuid=str(uuid4()), qr_secret=token_urlsafe(24))
        db.add(asset)
        created.append(asset)
        existing_codes.add(row.asset_code)
    db.commit()
    return {"data": {"created_count": len(created), "skipped": skipped}}


@router.post("", status_code=status.HTTP_201_CREATED)
def create_asset(payload: AssetCreate, _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    if db.scalar(select(Asset).where(Asset.asset_code == payload.asset_code)):
        raise HTTPException(status_code=409, detail="资产编号已存在")
    row = Asset(**payload.model_dump())
    db.add(row); db.commit(); db.refresh(row)
    return {"data": asset_payload(row)}


@router.patch("/{asset_id}")
def update_asset(asset_id: int, payload: AssetUpdate, _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(Asset, asset_id)
    if row is None: raise HTTPException(status_code=404, detail="资产不存在")
    for key, value in payload.model_dump(exclude_unset=True).items(): setattr(row, key, value)
    row.updated_at = datetime.now(UTC); db.commit(); db.refresh(row)
    return {"data": asset_payload(row)}


@router.delete("/{asset_id}")
def delete_asset(asset_id: int, _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(Asset, asset_id)
    if row is None: raise HTTPException(status_code=404, detail="资产不存在")
    db.delete(row); db.commit(); return {"data": {"id": asset_id, "deleted": True}}


@router.get("/{asset_id}/qr")
def get_asset_qr(asset_id: int, request: Request, _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(Asset, asset_id)
    if row is None:
        raise HTTPException(status_code=404, detail="资产不存在")
    return {"data": qr_payload(row, request)}


@router.post("/{asset_id}/qr/rotate")
def rotate_asset_qr(asset_id: int, request: Request, _: User = Depends(require_permission("asset:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(Asset, asset_id)
    if row is None:
        raise HTTPException(status_code=404, detail="资产不存在")
    row.qr_secret = token_urlsafe(24)
    db.commit()
    db.refresh(row)
    return {"data": qr_payload(row, request)}


@router.get("/mobile-inventory/asset")
def mobile_inventory_asset(asset_uuid: str, token: str, _: User = Depends(require_permission("inventory:self_check")), db: Session = Depends(get_db)) -> dict:
    asset = db.scalar(select(Asset).where(Asset.asset_uuid == asset_uuid))
    if asset is None or not hmac.compare_digest(asset_qr_token(asset), token):
        raise HTTPException(status_code=404, detail="资产二维码无效或已更新")
    active_record = active_record_payload(asset, db)
    return {
        "data": {
            "asset": asset_payload(asset),
            "active_record": active_record,
            "locked": bool(active_record and active_record["locked"]),
            "message": "暂无进行中的盘点任务" if active_record is None else "",
        }
    }


@router.post("/mobile-inventory/scan")
def submit_mobile_inventory(payload: MobileInventorySubmit, user: User = Depends(require_permission("inventory:self_check")), db: Session = Depends(get_db)) -> dict:
    asset = db.scalar(select(Asset).where(Asset.asset_uuid == payload.asset_uuid))
    if asset is None or not hmac.compare_digest(asset_qr_token(asset), payload.token):
        raise HTTPException(status_code=404, detail="资产二维码无效或已更新")
    task = db.get(InventoryTask, payload.task_id) if payload.task_id else None
    record = None
    if task:
        if task.status != "进行中":
            raise HTTPException(status_code=409, detail="本次盘点任务已结束，不能重复提交")
        record = db.scalar(select(InventoryRecord).where(InventoryRecord.task_id == task.id, InventoryRecord.asset_id == asset.id))
    if record is None:
        record = db.scalar(
            select(InventoryRecord)
            .join(InventoryTask, InventoryRecord.task_id == InventoryTask.id)
            .where(InventoryRecord.asset_id == asset.id, InventoryTask.status == "进行中")
            .order_by(InventoryRecord.id.desc())
        )
        task = db.get(InventoryTask, record.task_id) if record else task
    if record is None:
        raise HTTPException(status_code=409, detail="暂无进行中的盘点任务")
    if record.result != "未盘点":
        raise HTTPException(status_code=409, detail="本次盘点任务已完成，不能重复提交")
    remark_parts = []
    if payload.current_owner:
        remark_parts.append(f"现场使用人：{payload.current_owner}")
    if payload.remark:
        remark_parts.append(payload.remark)
    apply_inventory_record_update(
        record,
        asset,
        task,
        result=payload.result,
        actual_location=payload.actual_location,
        remark="\n".join(remark_parts) or None,
        photo_url=persist_photo_url(payload.photo_url, asset, record.id),
        checked_by=user.display_name,
        db=db,
    )
    db.commit()
    db.refresh(record)
    return {"data": record_payload(record, db)}


@router.get("/inventory-tasks")
def list_inventory_tasks(_: User = Depends(require_permission("asset:view")), db: Session = Depends(get_db)) -> dict:
    return {"data": [task_payload(row) for row in db.scalars(select(InventoryTask).order_by(InventoryTask.created_at.desc())).all()]}


@router.post("/inventory-tasks", status_code=status.HTTP_201_CREATED)
def create_inventory_task(payload: InventoryTaskCreate, _: User = Depends(require_permission("inventory:manage")), db: Session = Depends(get_db)) -> dict:
    count_query = select(func.count(Asset.id))
    scope_condition = inventory_scope_condition(payload.scope)
    if scope_condition is not None:
        count_query = count_query.where(scope_condition)
    total = db.scalar(count_query) or 0
    row = InventoryTask(**payload.model_dump(), total_count=total)
    db.add(row); db.commit(); db.refresh(row)
    return {"data": task_payload(row)}


@router.patch("/inventory-tasks/{task_id}")
def update_inventory_task(task_id: int, payload: InventoryTaskUpdate, _: User = Depends(require_permission("inventory:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(InventoryTask, task_id)
    if row is None: raise HTTPException(status_code=404, detail="盘点任务不存在")
    for key, value in payload.model_dump(exclude_unset=True).items(): setattr(row, key, value)
    row.updated_at = datetime.now(UTC); db.commit(); db.refresh(row)
    return {"data": task_payload(row)}


@router.post("/inventory-tasks/{task_id}/start")
def start_inventory_task(task_id: int, _: User = Depends(require_permission("inventory:manage")), db: Session = Depends(get_db)) -> dict:
    task = db.get(InventoryTask, task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="盘点任务不存在")
    if task.status in {"已结束", "已完成"}:
        raise HTTPException(status_code=409, detail="该盘点任务已结束，只能查看历史记录")
    for other in db.scalars(
        select(InventoryTask).where(InventoryTask.status == "进行中", InventoryTask.id != task.id)
    ).all():
        other.status = "已结束"
        other.updated_at = datetime.now(UTC)
    existing = db.scalar(select(func.count(InventoryRecord.id)).where(InventoryRecord.task_id == task_id)) or 0
    if not existing:
        asset_query = select(Asset).order_by(Asset.asset_code)
        scope_condition = inventory_scope_condition(task.scope)
        if scope_condition is not None:
            asset_query = asset_query.where(scope_condition)
        assets = db.scalars(asset_query).all()
        db.add_all([InventoryRecord(task_id=task_id, asset_id=asset.id) for asset in assets])
        task.total_count = len(assets)
    task.status = "进行中" if task.total_count else "已完成"
    task.updated_at = datetime.now(UTC)
    db.commit(); db.refresh(task)
    return {"data": task_payload(task)}


@router.get("/inventory-tasks/{task_id}/records")
def list_inventory_records(task_id: int, _: User = Depends(require_permission("asset:view")), db: Session = Depends(get_db)) -> dict:
    if db.get(InventoryTask, task_id) is None:
        raise HTTPException(status_code=404, detail="盘点任务不存在")
    rows = db.execute(
        select(InventoryRecord, Asset)
        .join(Asset, Asset.id == InventoryRecord.asset_id)
        .where(InventoryRecord.task_id == task_id)
        .order_by(InventoryRecord.id)
    ).all()
    return {"data": [record_payload(row, db, asset) for row, asset in rows]}


@router.patch("/inventory-records/{record_id}")
def update_inventory_record(record_id: int, payload: InventoryRecordUpdate, user: User = Depends(require_permission("inventory:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(InventoryRecord, record_id)
    if row is None:
        raise HTTPException(status_code=404, detail="盘点记录不存在")
    task = db.get(InventoryTask, row.task_id)
    asset = db.get(Asset, row.asset_id)
    apply_inventory_record_update(
        row,
        asset,
        task,
        result=payload.result,
        actual_location=payload.actual_location,
        remark=payload.remark,
        photo_url=persist_photo_url(payload.photo_url, asset, row.id),
        checked_by=user.display_name,
        db=db,
    )
    db.commit(); db.refresh(row)
    return {"data": record_payload(row, db)}


@router.get("/exceptions")
def list_exceptions(_: User = Depends(require_permission("asset:view")), db: Session = Depends(get_db)) -> dict:
    rows = db.execute(
        select(AssetException, Asset)
        .join(Asset, Asset.id == AssetException.asset_id)
        .order_by(AssetException.created_at.desc())
    ).all()
    return {"data": [exception_payload(row, db, asset) for row, asset in rows]}


@router.post("/exceptions/{exception_id}/resolve")
def resolve_exception(exception_id: int, payload: ResolveExceptionRequest, user: User = Depends(require_permission("inventory:manage")), db: Session = Depends(get_db)) -> dict:
    row = db.get(AssetException, exception_id)
    if row is None: raise HTTPException(status_code=404, detail="资产异常不存在")
    row.status = "已归档"; row.resolution = payload.resolution; row.resolved_by = user.display_name; row.resolved_at = datetime.now(UTC)
    db.commit(); db.refresh(row)
    return {"data": exception_payload(row, db)}
