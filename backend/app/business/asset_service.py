import csv
import hmac
import io
import socket
from base64 import b64decode, b64encode
from datetime import UTC, datetime
from hashlib import sha256
from uuid import uuid4

import qrcode
import qrcode.image.svg
from fastapi import HTTPException, Request, UploadFile
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.business.models import Asset, AssetException, InventoryRecord, InventoryTask
from app.core.config import settings
from app.integrations.storage import StorageUploadError, upload_bytes


def inventory_scope_condition(scope: str):
    """Translate the UI scope value into a database filter."""
    scope_filters = {
        "部门：": Asset.department,
        "类别：": Asset.category,
        "位置：": Asset.location,
    }
    for prefix, column in scope_filters.items():
        if scope.startswith(prefix):
            value = scope.removeprefix(prefix).strip()
            return column == value if value else None
    return None


def asset_payload(row: Asset) -> dict:
    return {"id": row.id, "asset_uuid": row.asset_uuid, "asset_code": row.asset_code, "name": row.name, "category": row.category, "department": row.department, "owner": row.owner, "location": row.location, "status": row.status, "purchase_date": row.purchase_date, "value": row.value, "last_inventory_at": row.last_inventory_at, "created_at": row.created_at, "updated_at": row.updated_at}


def task_payload(row: InventoryTask) -> dict:
    return {"id": row.id, "name": row.name, "scope": row.scope, "owner": row.owner, "status": row.status, "due_date": row.due_date, "total_count": row.total_count, "checked_count": row.checked_count, "exception_count": row.exception_count, "created_at": row.created_at, "updated_at": row.updated_at}


def record_payload(row: InventoryRecord, db: Session, asset: Asset | None = None) -> dict:
    asset = asset or db.get(Asset, row.asset_id)
    return {
        "id": row.id,
        "task_id": row.task_id,
        "asset_id": row.asset_id,
        "asset_code": asset.asset_code if asset else "",
        "asset_name": asset.name if asset else "",
        "department": asset.department if asset else "",
        "owner": asset.owner if asset else "",
        "ledger_location": asset.location if asset else "",
        "result": row.result,
        "actual_location": row.actual_location,
        "remark": row.remark,
        "photo_url": row.photo_url,
        "checked_by": row.checked_by,
        "checked_at": row.checked_at,
    }


def active_record_payload(asset: Asset, db: Session) -> dict | None:
    row = db.scalar(
        select(InventoryRecord)
        .join(InventoryTask, InventoryRecord.task_id == InventoryTask.id)
        .where(InventoryRecord.asset_id == asset.id, InventoryTask.status == "进行中")
        .order_by(InventoryRecord.id.desc())
    )
    if row is None:
        return None
    task = db.get(InventoryTask, row.task_id)
    locked = row.result != "未盘点"
    return {
        **record_payload(row, db),
        "task_name": task.name if task else "",
        "task_scope": task.scope if task else "",
        "task_due_date": task.due_date if task else None,
        "task_status": task.status if task else "",
        "locked": locked,
    }


def persist_photo_url(photo_url: str | None, asset: Asset, record_id: int) -> str | None:
    if not photo_url or not photo_url.startswith("data:"):
        return photo_url
    try:
        header, encoded = photo_url.split(",", 1)
        content_type = header[5:].split(";", 1)[0] or "image/jpeg"
        content = b64decode(encoded, validate=True)
    except (ValueError, UnicodeError):
        raise HTTPException(status_code=400, detail="现场照片格式无效") from None
    try:
        uploaded = upload_bytes(
            f"inventory/{asset.asset_uuid}/{record_id}-{uuid4().hex}",
            content,
            content_type,
        )
    except StorageUploadError as exc:
        if settings.app_env.lower() == "production":
            raise HTTPException(status_code=503, detail="云端照片存储不可用，请稍后重试") from exc
        return photo_url
    if uploaded:
        return uploaded
    if settings.app_env.lower() == "production":
        raise HTTPException(status_code=503, detail="生产环境必须配置云端照片存储")
    return photo_url


def exception_payload(row: AssetException, db: Session, asset: Asset | None = None) -> dict:
    asset = asset or db.get(Asset, row.asset_id)
    return {"id": row.id, "asset_id": row.asset_id, "asset_code": asset.asset_code if asset else "", "asset_name": asset.name if asset else "", "department": asset.department if asset else "", "type": row.type, "description": row.description, "status": row.status, "resolution": row.resolution, "resolved_by": row.resolved_by, "resolved_at": row.resolved_at, "created_at": row.created_at}


def normalize_import_key(value: str) -> str:
    aliases = {
        "资产编号": "asset_code",
        "固定ID": "asset_code",
        "资产名称": "name",
        "名称": "name",
        "类别": "category",
        "所属部门": "department",
        "部门": "department",
        "责任人": "owner",
        "位置": "location",
        "状态": "status",
        "资产价值": "value",
        "价值": "value",
    }
    key = value.strip()
    return aliases.get(key, key)


def parse_asset_rows(file: UploadFile) -> list[dict]:
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    content = file.file.read()
    if suffix == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [normalize_import_key(str(value or "")) for value in rows[0]]
        return [
            {header: value for header, value in zip(headers, row, strict=False)}
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = [normalize_import_key(value) for value in (reader.fieldnames or [])]
    return [row for row in reader]


def row_to_import(row: dict, row_number: int, existing_codes: set[str], seen_codes: set[str]) -> dict:
    required = ["asset_code", "name", "category", "department", "owner", "location"]
    cleaned = {key: str(row.get(key) or "").strip() for key in required}
    cleaned["status"] = str(row.get("status") or "在用").strip()
    try:
        cleaned["value"] = float(row.get("value") or 0)
    except (TypeError, ValueError):
        cleaned["value"] = -1
    errors = [f"缺少{key}" for key in required if not cleaned[key]]
    if cleaned["value"] < 0:
        errors.append("资产价值不能为负数")
    if cleaned["asset_code"] in existing_codes:
        errors.append("资产编号已存在")
    if cleaned["asset_code"] in seen_codes:
        errors.append("文件内资产编号重复")
    seen_codes.add(cleaned["asset_code"])
    return {"row_number": row_number, **cleaned, "valid": not errors, "errors": errors}


def asset_qr_token(asset: Asset) -> str:
    return hmac.new(
        settings.jwt_secret.encode(),
        f"{asset.asset_uuid}:{asset.qr_secret}".encode(),
        sha256,
    ).hexdigest()


def qr_image_data_url(value: str) -> str:
    image_factory = qrcode.image.svg.SvgPathImage
    image = qrcode.make(value, image_factory=image_factory)
    buffer = io.BytesIO()
    image.save(buffer)
    encoded = b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/svg+xml;base64,{encoded}"


def get_lan_ip() -> str:
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sock:
            sock.connect(("8.8.8.8", 80))
            return sock.getsockname()[0]
    except OSError:
        return "127.0.0.1"


def mobile_app_base_url(request: Request) -> tuple[str, str]:
    if settings.public_app_url:
        return settings.public_app_url.rstrip("/"), ""
    origin = request.headers.get("origin", "").rstrip("/")
    if origin and "127.0.0.1" not in origin and "localhost" not in origin:
        return origin, ""
    lan_ip = get_lan_ip()
    if lan_ip != "127.0.0.1":
        return (
            f"http://{lan_ip}:5173",
            "未配置 PUBLIC_APP_URL，已使用局域网地址。手机需与电脑在同一 WiFi，且前端服务需用 0.0.0.0 启动。",
        )
    return (
        "http://127.0.0.1:5173",
        "当前二维码仍是本机地址，微信手机无法打开。请配置 PUBLIC_APP_URL 为公网或局域网可访问地址。",
    )


def qr_payload(asset: Asset, request: Request) -> dict:
    token = asset_qr_token(asset)
    base_url, warning = mobile_app_base_url(request)
    mobile_url = f"{base_url}/assets/mobile-inventory?asset_uuid={asset.asset_uuid}&token={token}"
    return {
        "asset_id": asset.id,
        "asset_uuid": asset.asset_uuid,
        "asset_code": asset.asset_code,
        "token": token,
        "mobile_url": mobile_url,
        "mobile_url_warning": warning,
        "qr_image": qr_image_data_url(mobile_url),
    }


def refresh_task_progress(task: InventoryTask, db: Session, now: datetime) -> None:
    task.checked_count = db.scalar(select(func.count(InventoryRecord.id)).where(InventoryRecord.task_id == task.id, InventoryRecord.result != "未盘点")) or 0
    task.exception_count = db.scalar(
        select(func.count(AssetException.id))
        .join(InventoryRecord, AssetException.inventory_record_id == InventoryRecord.id)
        .where(InventoryRecord.task_id == task.id, AssetException.status != "已归档")
    ) or 0
    if task.status not in {"已结束", "已完成"}:
        task.status = "已完成" if task.total_count and task.checked_count >= task.total_count else "进行中"
    task.updated_at = now


def apply_inventory_record_update(
    record: InventoryRecord,
    asset: Asset | None,
    task: InventoryTask | None,
    *,
    result: str,
    actual_location: str | None,
    remark: str | None,
    photo_url: str | None,
    checked_by: str,
    db: Session,
) -> None:
    now = datetime.now(UTC)
    record.result = result
    record.actual_location = actual_location or (asset.location if asset else None)
    record.remark = remark
    record.photo_url = photo_url
    record.checked_by = checked_by
    record.checked_at = now
    if asset:
        asset.last_inventory_at = now.date()

    existing_exception = db.scalar(select(AssetException).where(AssetException.inventory_record_id == record.id))
    if result == "正常":
        if existing_exception and existing_exception.status != "已归档":
            existing_exception.status = "已归档"
            existing_exception.resolution = "盘点结果已更正为正常"
            existing_exception.resolved_by = checked_by
            existing_exception.resolved_at = now
    elif existing_exception:
        existing_exception.type = result
        existing_exception.description = remark or f"盘点结果：{result}"
        existing_exception.status = "待核实"
    else:
        db.add(AssetException(asset_id=record.asset_id, inventory_record_id=record.id, type=result, description=remark or f"盘点结果：{result}"))

    db.flush()
    if task:
        refresh_task_progress(task, db, now)
